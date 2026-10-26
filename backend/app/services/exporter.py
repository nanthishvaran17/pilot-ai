import os
import json
import csv
from datetime import datetime
from typing import Tuple
from sqlalchemy.orm import Session
from app.db.models import Project, Blueprint, Requirement, ExportJob
from app.core.config import settings

# ReportLab imports for real PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# python-docx imports for real DOCX generation
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

class NumberedCanvas(canvas.Canvas):
    """Adds 'Page X of Y' and header to ReportLab documents"""
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 750, "ProjectPilot AI — Project Blueprint & Specification")
            self.setStrokeColor(colors.HexColor("#e2e8f0"))
            self.setLineWidth(0.5)
            self.line(54, 742, letter[0] - 54, 742)

        # Footer
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 36, footer_text)
        self.drawString(54, 36, "Confidential — Generated for Engineering Implementation")
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        self.line(54, 48, letter[0] - 54, 48)
        self.restoreState()


class ExportEngine:
    @classmethod
    def generate_export(cls, project: Project, export_format: str, db: Session) -> ExportJob:
        bp = db.query(Blueprint).filter(Blueprint.project_id == project.id).first()
        if not bp:
            from app.services.blueprint_generator import BlueprintGenerator
            bp = BlueprintGenerator.generate(project, db)

        format_lower = export_format.lower()
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        safe_name = "".join(c for c in project.name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
        if not safe_name:
            safe_name = "project"

        file_name = f"{safe_name}_blueprint_{timestamp}.{format_lower}"
        export_dir = os.path.join(settings.STORAGE_DIR, "exports")
        os.makedirs(export_dir, exist_ok=True)
        file_path = os.path.join(export_dir, file_name)

        if format_lower == "pdf":
            cls._create_pdf(project, bp, file_path)
        elif format_lower == "docx":
            cls._create_docx(project, bp, file_path)
        elif format_lower == "json":
            cls._create_json(project, bp, file_path)
        elif format_lower == "csv":
            cls._create_csv(project, db, file_path)
        elif format_lower in ["md", "markdown"]:
            cls._create_markdown(project, bp, file_path)
            format_lower = "md"
        else:
            raise ValueError(f"Unsupported export format: {export_format}")

        file_size = os.path.getsize(file_path) if os.path.exists(file_path) else 0

        job = ExportJob(
            project_id=project.id,
            export_format=format_lower,
            file_path=file_path,
            file_name=file_name,
            file_size=file_size,
            status="completed"
        )
        db.add(job)
        db.commit()
        db.refresh(job)
        return job

    @staticmethod
    def _create_markdown(project: Project, bp: Blueprint, file_path: str):
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(bp.markdown_content)

    @staticmethod
    def _create_json(project: Project, bp: Blueprint, file_path: str):
        payload = {
            "project_id": project.id,
            "project_name": project.name,
            "version": bp.version,
            "created_at": bp.created_at.isoformat(),
            "blueprint": bp.content_json
        }
        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2)

    @staticmethod
    def _create_csv(project: Project, db: Session, file_path: str):
        reqs = db.query(Requirement).filter(Requirement.project_id == project.id).all()
        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["ID", "Title", "Category", "MoSCoW Priority", "Importance (1-5)", "Complexity", "Time Savings", "Target Users", "MVP", "Source"])
            for r in reqs:
                writer.writerow([
                    r.id,
                    r.title,
                    r.category,
                    r.moscow_priority.replace("_", " ").title(),
                    r.importance,
                    r.complexity,
                    r.time_savings or "N/A",
                    r.users_affected or "All",
                    "Yes" if r.is_mvp else "No",
                    r.source
                ])

    @staticmethod
    def _create_docx(project: Project, bp: Blueprint, file_path: str):
        doc = Document()
        
        # Styles
        title = doc.add_heading(project.name, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        sub = doc.add_paragraph("AI-Engineered Requirements & Architecture Blueprint")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.style.font.color.rgb = RGBColor(100, 116, 139)
        
        doc.add_paragraph() # spacer

        data = bp.content_json

        # 1. Executive Summary
        doc.add_heading("1. Executive Summary", level=1)
        doc.add_paragraph(data.get("executive_summary", ""))

        # 2. Problem Statement
        doc.add_heading("2. Problem Statement", level=1)
        doc.add_paragraph(data.get("problem_statement", ""))

        # 3. Project Goals
        doc.add_heading("3. Project Goals", level=1)
        for g in data.get("project_goals", []):
            doc.add_paragraph(g, style='List Bullet')

        # 4. Target Users & Roles
        doc.add_heading("4. Target Users & Roles", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Role"
        hdr_cells[1].text = "Persona"
        hdr_cells[2].text = "Primary Needs"
        for u in data.get("target_users", []):
            row_cells = table.add_row().cells
            row_cells[0].text = str(u.get("role", ""))
            row_cells[1].text = str(u.get("persona", ""))
            row_cells[2].text = str(u.get("primary_needs", ""))

        doc.add_paragraph()

        # 5. Confirmed Functional Requirements
        doc.add_heading("5. Confirmed Functional Requirements", level=1)
        t_req = doc.add_table(rows=1, cols=4)
        t_req.style = 'Table Grid'
        h = t_req.rows[0].cells
        h[0].text = "Title"
        h[1].text = "Category"
        h[2].text = "Priority"
        h[3].text = "MVP"
        for r in data.get("confirmed_functional_requirements", []):
            rc = t_req.add_row().cells
            rc[0].text = str(r.get("title", ""))
            rc[1].text = str(r.get("category", ""))
            rc[2].text = str(r.get("priority", ""))
            rc[3].text = "Yes" if r.get("is_mvp") else "No"

        doc.add_paragraph()

        # 6. Page-by-page UI Plan
        doc.add_heading("6. Page-by-Page UI Plan", level=1)
        for p in data.get("page_by_page_ui_plan", []):
            doc.add_paragraph(f"{p.get('page')} ({p.get('route')})", style='List Bullet')
            doc.add_paragraph(f"Components: {p.get('components')}")

        # 7. AI Architecture
        doc.add_heading("7. AI Architecture & Boundaries", level=1)
        ai = data.get("ai_architecture", {})
        doc.add_paragraph(f"Strategy: {ai.get('provider_strategy')}")
        doc.add_paragraph(f"Deterministic Rules: {ai.get('deterministic_boundary')}")
        doc.add_paragraph(f"Fault Tolerance: {ai.get('fallback_mechanism')}")

        # 8. Security & Compliance
        doc.add_heading("8. Security & Compliance", level=1)
        for s in data.get("security_requirements", []):
            doc.add_paragraph(s, style='List Bullet')

        # 9. Acceptance Criteria
        doc.add_heading("9. Acceptance Criteria", level=1)
        for ac in data.get("acceptance_criteria", []):
            doc.add_paragraph(f"[ ] {ac}")

        doc.save(file_path)

    @staticmethod
    def _create_pdf(project: Project, bp: Blueprint, file_path: str):
        doc = SimpleDocTemplate(
            file_path,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        
        # Custom palette & typography
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=6
        )

        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#475569"),
            spaceAfter=15
        )

        h1_style = ParagraphStyle(
            'H1',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=14,
            spaceAfter=6
        )

        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6
        )

        bullet_style = ParagraphStyle(
            'Bullet',
            parent=body_style,
            leftIndent=15,
            firstLineIndent=-10,
            spaceAfter=3
        )

        table_header_style = ParagraphStyle(
            'TableHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=11,
            textColor=colors.white
        )

        table_cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#1e293b")
        )

        story = []
        data = bp.content_json

        # Title Block
        story.append(Paragraph(project.name, title_style))
        story.append(Paragraph("Project Requirements & Architectural Specification Blueprint", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#3b82f6"), spaceAfter=14))

        # Executive Summary
        story.append(Paragraph("1. Executive Summary", h1_style))
        story.append(Paragraph(data.get("executive_summary", ""), body_style))

        # Problem Statement
        story.append(Paragraph("2. Problem Statement", h1_style))
        story.append(Paragraph(data.get("problem_statement", ""), body_style))

        # Project Goals
        story.append(Paragraph("3. Project Goals", h1_style))
        for g in data.get("project_goals", []):
            story.append(Paragraph(f"• {g}", bullet_style))

        # Confirmed Requirements Table
        story.append(Paragraph("4. Confirmed Requirements", h1_style))
        table_rows = [[
            Paragraph("Requirement", table_header_style),
            Paragraph("Category", table_header_style),
            Paragraph("Priority", table_header_style),
            Paragraph("MVP", table_header_style)
        ]]
        for r in data.get("confirmed_functional_requirements", [])[:10]:
            table_rows.append([
                Paragraph(str(r.get("title", "")), table_cell_style),
                Paragraph(str(r.get("category", "")), table_cell_style),
                Paragraph(str(r.get("priority", "")), table_cell_style),
                Paragraph("Yes" if r.get("is_mvp") else "No", table_cell_style)
            ])

        t = Table(table_rows, colWidths=[240, 100, 100, 50])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))

        # UI Pages
        story.append(Paragraph("5. Page-by-Page UI Architecture", h1_style))
        for p in data.get("page_by_page_ui_plan", []):
            story.append(Paragraph(f"<b>{p.get('page')}</b> — <i>Route: {p.get('route')}</i>", body_style))
            story.append(Paragraph(f"Key Components: {p.get('components')}", bullet_style))

        # Security & Compliance
        story.append(Paragraph("6. Security, Compliance & Governance", h1_style))
        for s in data.get("security_requirements", []):
            story.append(Paragraph(f"• {s}", bullet_style))

        # Acceptance Criteria
        story.append(Paragraph("7. Acceptance Criteria", h1_style))
        for ac in data.get("acceptance_criteria", []):
            story.append(Paragraph(f"• [ ] {ac}", bullet_style))

        doc.build(story, canvasmaker=NumberedCanvas)
